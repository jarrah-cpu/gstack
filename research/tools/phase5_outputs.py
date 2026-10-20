#!/usr/bin/env python3
"""Phase 5 outputs from data/events.jsonl:

  output/events_master.csv    one row per event
  output/events_master.xlsx   All events, then tabs by category and by region; filters on, links clickable
  output/calendar.md          next editions from today to the 18-month horizon, chronological
  output/brand_map.md         brands and the events they sponsor, with role, years and source
  output/shortlist_data.json  top 15 overall and top 5 per category by weighted_total (input for shortlist.md)

Usage: python3 research/tools/phase5_outputs.py
"""
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
TODAY, HORIZON = "2026-10-21", "2028-04-20"
events = [json.loads(l) for l in (ROOT / "data" / "events.jsonl").read_text().split("\n") if l.strip()]

CATEGORY_LABEL = {
    "streetwear-culture-festival": "Streetwear culture", "creative-market": "Creative market",
    "design-art-week": "Design and art week", "skate-surf-culture-weekend": "Skate and surf",
    "independent-music-festival": "Independent music", "club-culture": "Club culture",
    "makers-vintage": "Makers and vintage", "food-culture-night-market": "Food and night market",
    "subculture-festival": "Subculture festival", "competition-or-contest": "Competition or contest",
    "expo-or-trade-show": "Expo or trade show", "other": "Other"}


def sponsors_str(e):
    return "; ".join(f"{s['brand']} ({s.get('role') or 'unknown'}{', ' + ', '.join(map(str, s['years'])) if s['years'] else ''})"
                     for s in e["sponsors"])


def row(e):
    d, ig, at, au, org = e["dates"], e["instagram"], e["attendance"], e["audience"], e["organiser"]
    pc, tp, v, sc = e["partnership_contact"], e["ticket_price"], e["vibe"], e["scores"]
    return {
        "id": e["id"], "name": e["name"], "category": CATEGORY_LABEL.get(e["category"], e["category"]),
        "subculture_tags": ", ".join(e["subculture_tags"]), "scope": e["scope"], "region": e["region"],
        "city": e["location"]["city"], "state_or_region": e["location"]["state_or_region"],
        "country": e["location"]["country"], "venue_or_area": e["location"]["venue_or_area"], "status": e["status"],
        "last_edition_start": d["last_edition"]["start"], "last_edition_end": d["last_edition"]["end"],
        "next_edition_start": d["next_edition"]["start"], "next_edition_end": d["next_edition"]["end"],
        "next_edition_status": d["next_edition"]["status"], "frequency": d["frequency"], "duration_days": d["duration_days"],
        "website": e["website"], "ig_handle": ig["handle"], "ig_url": ig["url"], "ig_followers": ig["followers"],
        "ig_followers_as_of": ig["followers_as_of"], "ig_post_count": ig.get("post_count"),
        "ig_avg_recent_engagement": ig.get("avg_recent_engagement"), "organiser_ig_handle": ig["organiser_handle"],
        "organiser_ig_url": ig["organiser_url"], "organiser_ig_followers": ig["organiser_followers"],
        "attendance": at["figure"], "attendance_high": at.get("figure_high"), "attendance_year": at["year"],
        "attendance_is_estimate": at["is_estimate"], "attendance_source": at["source"], "attendance_note": at["note"],
        "audience_age_range": au["age_range"], "audience_gender_skew": au["gender_skew"],
        "audience_psychographics": au["psychographics"], "audience_source": au["source"],
        "organiser": org["name"], "organiser_type": org.get("type"), "organiser_website": org["website"],
        "sponsors": sponsors_str(e), "sponsor_count": len(e["sponsors"]),
        "activation_formats": ", ".join(e["activation_formats"]),
        "partnership_email": pc["email"], "partnership_page": pc["page_url"], "partnership_person": pc["named_person"],
        "partnership_source": pc["source"], "ticket_min": tp["min"], "ticket_max": tp["max"], "ticket_currency": tp["currency"],
        "free_entry": tp["free_entry"], "ticket_note": tp["note"], "vibe_score": v["score"], "vibe_rationale": v["rationale"],
        "vibe_preliminary": v["preliminary"], "vibe_fit": sc["vibe_fit"], "audience_fit": sc["audience_fit"],
        "scale_band": sc["scale_band"], "accessibility": sc["accessibility"], "cost_signal": sc["cost_signal"],
        "weighted_total": sc["weighted_total"], "fit_notes": e["fit_notes"], "phase": e["record_meta"]["phase"],
        "verifier_flags": " | ".join(e["record_meta"].get("verifier_flags") or []), "gaps": " | ".join(e["gaps"]),
        "source_count": len(e["sources"]), "source_urls": " ; ".join(dict.fromkeys(s["url"] for s in e["sources"])),
    }


rows = [row(e) for e in sorted(events, key=lambda e: (-(e["scores"]["weighted_total"] or 0), e["name"].lower()))]
cols = list(rows[0])
with (OUT / "events_master.csv").open("w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)

# Excel workbook
LINKS = {"website", "ig_url", "organiser_ig_url", "attendance_source", "partnership_page", "partnership_source", "organiser_website"}
WIDTH = {"name": 34, "fit_notes": 60, "vibe_rationale": 50, "sponsors": 50, "audience_psychographics": 50, "gaps": 60,
         "source_urls": 60, "attendance_note": 40, "ticket_note": 30, "verifier_flags": 50}


def sheet(wb, title, rs):
    ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", title)[:31])
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4F4F")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rs:
        ws.append([r[c] for c in cols])
        n = ws.max_row
        for i, c in enumerate(cols, 1):
            if c in LINKS and r[c] and str(r[c]).startswith("http"):
                cell = ws.cell(n, i)
                cell.hyperlink = r[c]
                cell.font = Font(color="0563C1", underline="single")
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(c, 16)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


wb = Workbook()
wb.remove(wb.active)
sheet(wb, "All events", rows)
by_cat, by_reg = defaultdict(list), defaultdict(list)
for r in rows:
    by_cat[r["category"]].append(r)
    for reg in r["region"].split("+"):
        by_reg[reg].append(r)
for k in sorted(by_cat):
    sheet(wb, f"Cat {k}", by_cat[k])
for k in sorted(by_reg):
    sheet(wb, f"Reg {k}", by_reg[k])
wb.save(OUT / "events_master.xlsx")

# calendar
def start_key(s):
    return (s + "-99")[:10] if s else "9999"


cal = [e for e in events if e["status"] in ("active", "unknown") and e["dates"]["next_edition"]["start"]
       and e["dates"]["next_edition"]["status"] in ("confirmed", "expected")
       and TODAY[:len(e["dates"]["next_edition"]["start"])] <= e["dates"]["next_edition"]["start"] <= HORIZON]
cal.sort(key=lambda e: (start_key(e["dates"]["next_edition"]["start"]), e["name"].lower()))
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def when(n):
    s, en = n["start"], n["end"]
    def fmt(x):
        if not x:
            return None
        p = x.split("-")
        if len(p) == 1:
            return p[0]
        if len(p) == 2:
            return f"{MONTHS[int(p[1]) - 1][:3]} {p[0]}"
        return f"{int(p[2])} {MONTHS[int(p[1]) - 1][:3]} {p[0]}"
    return fmt(s) + (f" to {fmt(en)}" if en and en != s else "")


L = ["# Event calendar: next editions, 21 Oct 2026 to 20 Apr 2028", "",
     f"Generated from `data/events.jsonl` by `research/tools/phase5_outputs.py`. {len(cal)} events with a confirmed or expected next edition in the window, in date order. "
     "Expected dates are inferred from past patterns and are not yet announced; month-only or year-only dates sort to the end of their month or year. "
     "Score is the Phase 4 weighted total (1 to 5).", ""]
cur = None
for e in cal:
    s = e["dates"]["next_edition"]["start"]
    p = s.split("-")
    head = f"{MONTHS[int(p[1]) - 1]} {p[0]}" if len(p) >= 2 else f"{p[0]} (month not known)"
    if head != cur:
        L += ["", f"## {head}", "", "| When | Event | City | Category | Status | Score | Instagram |", "|---|---|---|---|---|---|---|"]
        cur = head
    ig = e["instagram"]["handle"] or e["instagram"]["organiser_handle"] or ""
    igl = f"[{ig}](https://www.instagram.com/{ig[1:]}/)" if ig else ""
    L.append(f"| {when(e['dates']['next_edition'])} | {e['name']} | {e['location']['city'] or ''}, {e['location']['country']} | "
             f"{CATEGORY_LABEL.get(e['category'])} | {e['dates']['next_edition']['status']} | {e['scores']['weighted_total'] or ''} | {igl} |")
(OUT / "calendar.md").write_text("\n".join(L) + "\n")

# brand map
bm = defaultdict(list)
for e in events:
    for s in e["sponsors"]:
        bm[s["brand"].strip()].append((e, s))
brands = sorted(bm.items(), key=lambda kv: (-len(kv[1]), kv[0].lower()))
L = ["# Brand map: who sponsors what", "",
     f"Generated from `data/events.jsonl`. {len(brands)} brands and organisations across {sum(1 for e in events if e['sponsors'])} events. "
     "Role and years are as recorded; each row links its source. Brands with two or more events come first.", "",
     "## Brands at two or more events", "", "| Brand | Events | Event, role, years |", "|---|---|---|"]
for b, lst in brands:
    if len(lst) < 2:
        continue
    cells = "; ".join(f"[{e['name']}]({s['source']}) ({s.get('role') or 'unknown'}{', ' + ', '.join(map(str, s['years'])) if s['years'] else ''})"
                      if s.get("source") else f"{e['name']} ({s.get('role') or 'unknown'})" for e, s in lst)
    L.append(f"| {b} | {len(lst)} | {cells} |")
L += ["", "## All other brands (one event each)", "", "| Brand | Event | Role | Years | Source |", "|---|---|---|---|---|"]
for b, lst in brands:
    if len(lst) >= 2:
        continue
    e, s = lst[0]
    L.append(f"| {b} | {e['name']} | {s.get('role') or 'unknown'} | {', '.join(map(str, s['years']))} | {s.get('source') or ''} |")
(OUT / "brand_map.md").write_text("\n".join(L) + "\n")

# shortlist data
scored = [e for e in events if e["scores"]["weighted_total"] is not None and e["status"] in ("active", "unknown")]
scored.sort(key=lambda e: -e["scores"]["weighted_total"])
short = {"top15": [e["id"] for e in scored[:15]], "by_category": {}}
for cat in CATEGORY_LABEL:
    short["by_category"][cat] = [e["id"] for e in scored if e["category"] == cat][:5]
(OUT / "shortlist_data.json").write_text(json.dumps(short, indent=1))
print(f"{len(rows)} rows; calendar {len(cal)}; brands {len(brands)}; scored {len(scored)}")
